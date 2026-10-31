"""Make the homepage's PowerPoint demonstration with the real application.

Runs the published 2.3.0 source (see docs/CAPTURE.md) on the native macOS platform. It opens the
sample blade drawing in the real MARK window, adds three numbered callouts with linked comments, a
circle, two arrows and a highlight, sends the page to the PowerPoint queue and exports the queue
with the app's own export controller, so the .pptx is exactly what the app writes.

    <app>/.venv/bin/python tools/demo/make_review_slide.py --snapshot /path/to/snapshot --site .

Writes public/demo/review/blade-review.pptx, a WebP of the picture embedded in the deck, and
src/data/review.json, which holds the same marks in the sheet's coordinates for the web animations.
A window opens for about half a minute; settings and caches go to a temporary folder.
"""
import argparse
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import time

# Everything below is in PDF points, which match the sheet's SVG viewBox units (820 x 540).
BUBBLE = 34
COMMENTS = ['Review the midspan section.', 'Review the blade root fit.', 'Check the airfoil surface.']
RING = (596, 116, 712, 178)
BUBBLES = [(744, 118), (532, 392), (330, 118)]
ARROWS = [((515, 389), (466, 384)), ((313, 123), (236, 146))]
HIGHLIGHT = ((226, 130), (231, 208))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--snapshot', type=Path, required=True)
    parser.add_argument('--site', type=Path, required=True)
    args = parser.parse_args()
    site = args.site.resolve()
    sys.path.insert(0, str(args.snapshot.resolve() / 'src'))
    os.environ.pop('QT_QPA_PLATFORM', None)
    os.environ.pop('QT_SCALE_FACTOR', None)

    from PySide6.QtCore import QPointF, QRectF
    from PySide6.QtGui import QColor, QPixmap
    from PySide6.QtWidgets import QApplication
    import eng_viewer
    import eng_viewer.core.config as config
    import eng_viewer.ui.thumbnails as thumbnails
    from eng_viewer.core.config import Settings
    from eng_viewer.core.render_service import RenderService
    from eng_viewer.ui.export_dialog import ExportController
    from eng_viewer.ui.main_window import MainWindow
    from eng_viewer.ui.mark.commands import AddItems
    from eng_viewer.ui.mark.items import ArrowItem, BubbleItem, EllipseItem, StrokeItem, BUBBLE_RED
    from eng_viewer.ui.theme import ThemeManager

    work = Path(tempfile.mkdtemp(prefix='eng-site-review-'))
    os.chdir(site / 'design' / 'drawings')
    blade = Path('Demo Project') / 'Turbine Stage' / 'ENG-003 Rotor blade.pdf'
    out = site / 'public' / 'demo' / 'review'
    out.mkdir(parents=True, exist_ok=True)

    app = QApplication.instance() or QApplication([])
    app.setApplicationName('ENG Drawing Viewer')
    config.cache_dir = lambda: work
    thumbnails.cache_dir = lambda: work
    settings = Settings(work / 'settings.json')
    settings.set('recent_folders', [])
    settings.set('thumbnail_disk_cache', False)
    tm = ThemeManager(app, 'light')
    tm.apply()
    render = RenderService(max_workers=3)
    w = MainWindow(settings, render, QPixmap())
    w.resize(1440, 900)
    w.show()

    def wait(seconds=1.0):
        until = time.monotonic() + seconds
        while time.monotonic() < until:
            app.processEvents()
            time.sleep(0.01)

    w.load_folder(Path('Demo Project'))
    wait(2)
    w.open_document(str(blade), 0)
    wait(2)
    w.act_mark.trigger()
    wait(3)
    mark = next(iter(w._mark_windows.values()))
    mark.resize(1440, 900)
    mark.raise_()
    mark.canvas.fit_page()
    wait(1.5)
    canvas = mark.canvas

    x0, y0, x1, y1 = RING
    canvas.undo_stack.push(AddItems(canvas, [EllipseItem(QRectF(x0, y0, x1 - x0, y1 - y0), BUBBLE_RED, 2.6)], 'Circle'))
    for (a, b) in ARROWS:
        canvas.undo_stack.push(AddItems(canvas, [ArrowItem(QPointF(*a), QPointF(*b), BUBBLE_RED, 2.6)], 'Arrow'))
    highlight = StrokeItem(QColor('#FFD21E'), 15, highlighter=True, opacity=0.42)
    steps = 24
    for i in range(steps + 1):
        t = i / steps
        highlight.add_point(QPointF(HIGHLIGHT[0][0] + (HIGHLIGHT[1][0] - HIGHLIGHT[0][0]) * t,
                                    HIGHLIGHT[0][1] + (HIGHLIGHT[1][1] - HIGHLIGHT[0][1]) * t))
    canvas.undo_stack.push(AddItems(canvas, [highlight], 'Highlight'))
    for i, ((x, y), text) in enumerate(zip(BUBBLES, COMMENTS), 1):
        bubble = BubbleItem(i, BUBBLE, True, canvas.font().family())
        bubble.setPos(QPointF(x, y))
        canvas.undo_stack.push(AddItems(canvas, [bubble], f'Add callout {i}'))
        wait(0.1)
        mark.comments.rows[id(bubble)].edit.setPlainText(text)
    canvas.set_tool('select')
    wait(1)

    mark.send_to_queue()
    wait(2)
    entries = list(w.queue.model.entries)
    assert len(entries) == 1 and entries[0].marked, 'the marked page did not reach the queue'
    pptx = out / 'blade-review.pptx'
    done = {}
    controller = ExportController(render, entries, {
        'output': str(pptx), 'mode': 'new', 'slide_size': '16:9', 'title_mode': 'title', 'quality': 'high'}, w)
    controller.finished.connect(lambda info: done.update(info))
    controller.failed.connect(lambda message: done.update(error=message))
    controller.start()
    until = time.monotonic() + 60
    while not done and time.monotonic() < until:
        app.processEvents()
        time.sleep(0.02)
    print('export:', done, flush=True)
    assert done and 'error' not in done, 'the export failed'
    canvas.undo_stack.setClean()

    from pptx import Presentation
    prs = Presentation(str(pptx))
    slide = prs.slides[0]
    picture = next(s for s in slide.shapes if s.shape_type == 13)
    from PIL import Image
    marked = Image.open(io.BytesIO(picture.image.blob)).convert('RGB')
    marked.thumbnail((1800, 1800), Image.Resampling.LANCZOS)
    marked.save(out / 'blade-marked.webp', 'WEBP', quality=90, method=6)
    boxes = [s for s in slide.shapes if s.has_text_frame]
    title = slide.shapes.title.text_frame.text if slide.shapes.title is not None else ''
    comment_box = next(s for s in boxes if s != slide.shapes.title)
    comments = [p.text for p in comment_box.text_frame.paragraphs if p.text.strip()]
    data = {
        'app_version': eng_viewer.__version__,
        'sheet': {'width': 820, 'height': 540},
        'bubble_diameter': BUBBLE,
        'ring': RING,
        'bubbles': [{'number': i, 'x': x, 'y': y, 'comment': c}
                    for i, ((x, y), c) in enumerate(zip(BUBBLES, COMMENTS), 1)],
        'arrows': ARROWS,
        'highlight': HIGHLIGHT,
        'slide': {'title': title, 'comments': comments, 'picture_px': [picture.image.size[0], picture.image.size[1]],
                  'width_emu': prs.slide_width, 'height_emu': prs.slide_height,
                  'picture_box': [picture.left, picture.top, picture.width, picture.height],
                  'comment_box': [comment_box.left, comment_box.top, comment_box.width, comment_box.height]},
        'pptx_bytes': pptx.stat().st_size,
    }
    (site / 'src' / 'data' / 'review.json').write_text(json.dumps(data, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(data['slide'], indent=1), flush=True)
    render.shutdown()
    os._exit(0)


if __name__ == '__main__':
    main()
